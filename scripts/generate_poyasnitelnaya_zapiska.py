#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Генерация пояснительной записки к курсовому проекту (Word .docx).

Ориентир по оформлению: методические указания (Times New Roman 14 pt,
одинарный интервал, выравнивание по ширине, абзац 1,25 см, поля А4:
слева 23 мм, справа 10 мм, сверху 20 мм, снизу 15 мм).

Дальше вручную в Word: титульный лист и задание без номеров страниц;
колонтитулы с нумерацией с введения; автоматическое оглавление по стилям
или таблица по приложению А; скриншоты интерфейса; согласование с
руководителем.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Mm, Pt


def set_run_tnr(run, size_pt=14, bold=False):
    run.font.name = 'Times New Roman'
    run.bold = bold
    run.font.size = Pt(size_pt)
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn('w:ascii'), 'Times New Roman')
    rfonts.set(qn('w:hAnsi'), 'Times New Roman')
    rfonts.set(qn('w:cs'), 'Times New Roman')
    rfonts.set(qn('w:eastAsia'), 'Times New Roman')


def configure_section(document):
    sec = document.sections[0]
    sec.page_height = Mm(297)
    sec.page_width = Mm(210)
    sec.left_margin = Mm(23)
    sec.right_margin = Mm(10)
    sec.top_margin = Mm(20)
    sec.bottom_margin = Mm(15)


def configure_normal_style(document):
    st = document.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(14)
    pf = st.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)


def add_center_line(document, text, bold=False, space_after=6):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_tnr(run, bold=bold)


def add_body(document, text):
    p = document.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Mm(12.5)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_tnr(run, bold=False)
    return p


def add_section_heading(document, text):
    p = document.add_paragraph()
    p.paragraph_format.first_line_indent = Mm(12.5)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_tnr(run, bold=True)
    return p


def add_subsection_heading(document, text):
    return add_section_heading(document, text)


def add_center_major(document, text):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_tnr(run, bold=True)
    return p


def build_title_page(document, args):
    add_center_line(document, args.ministry, False, 4)
    add_center_line(document, args.university, True, 4)
    add_center_line(document, 'Факультет: ' + args.faculty, False, 2)
    add_center_line(document, 'Кафедра: ' + args.department, False, 2)
    add_center_line(document, 'Специальность / направление: ' + args.specialty, False, 8)

    add_center_line(document, 'ПОЯСНИТЕЛЬНАЯ ЗАПИСКА', True, 4)
    add_center_line(document, 'к курсовому проекту', False, 4)
    add_center_line(document, 'по дисциплине «' + args.discipline + '»', False, 8)
    add_center_line(document, 'Тема: ' + args.topic, True, 16)

    add_center_line(document, 'Исполнитель', False, 2)
    add_center_line(document, args.student, False, 2)
    add_center_line(document, args.group_line, False, 12)

    add_center_line(document, 'Руководитель', False, 2)
    add_center_line(document, args.supervisor, False, 8)
    add_center_line(document, args.supervisor_title, False, 12)

    add_center_line(document, '«____» _____________ 20___ г.', False, 4)
    add_center_line(document, 'подпись, дата', False, 16)


def build_assignment_stub(document):
    add_center_major(document, 'ЗАДАНИЕ НА КУРСОВОЕ ПРОЕКТИРОВАНИЕ')
    add_body(
        document,
        'Вставьте сюда скан или набор текста официального задания '
        '(подписи студента, руководителя, заведующего кафедрой). '
        'Название темы в задании и на титульном листе должно совпадать.',
    )
    add_body(
        document,
        'Объём пояснительной записки уточняйте у руководителя '
        '(ориентир до 25–30 страниц с иллюстрациями, без учёта приложений).',
    )


def build_toc_stub(document):
    add_center_major(document, 'СОДЕРЖАНИЕ')
    add_body(
        document,
        'Оформите содержание по приложению А к методическим указаниям: '
        'таблица из двух столбцов, отточие до номера страницы, выравнивание '
        'номеров по разрядам. В Word удобно использовать '
        'Ссылки → Оглавление после назначения стилей заголовков.',
    )
    lines = [
        'Введение',
        '1 Постановка задачи',
        '1.1 Анализ предметной области и аналогов',
        '1.2 Техническое задание',
        '1.3 Выбор средств реализации',
        '1.4 Выводы по разделу',
        '2 Проектирование веб-приложения',
        '2.1 Структура страниц и навигация',
        '2.2 Макет интерфейса и компоненты',
        '2.3 Выводы по разделу',
        '3 Реализация',
        '3.1 Средства разработки и структура проекта',
        '3.2 Загрузка и разбор данных в формате XML',
        '3.3 Интерфейс чата и сохранение состояния',
        '3.4 Выводы по разделу',
        '4 Тестирование и руководство пользователя',
        '4.1 Проверка сценариев',
        '4.2 Руководство пользователя',
        '4.3 Выводы по разделу',
        'Заключение',
        'Список использованных источников',
        'Приложения (скриншоты, листинги, XML)',
    ]
    for line in lines:
        p = document.add_paragraph()
        p.paragraph_format.first_line_indent = Mm(0)
        p.paragraph_format.left_indent = Mm(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(line + '\t[стр.]')
        set_run_tnr(run, bold=False)


def build_intro(document, args):
    add_center_major(document, 'ВВЕДЕНИЕ')
    add_body(
        document,
        'Студенты при подготовке к занятиям и экзаменам часто обращаются '
        'к разрозненным материалам: конспектам, статьям, документации. '
        'Наглядный тематический справочник в виде веб-приложения сокращает '
        'время поиска базовых определений и ссылок на авторитетные источники.',
    )
    add_body(
        document,
        'Актуальность темы связана с распространением одностраничных '
        'приложений на JavaScript и необходимостью отделить данные '
        '(темы помощи) от кода, чтобы их мог обновлять преподаватель или '
        'методист без пересборки логики.',
    )
    add_body(
        document,
        'Цель курсового проекта — разработать веб-приложение '
        '«' + args.topic + '» с выбором учебной темы, кратким описанием, '
        'ссылкой на материал и диалоговым окном для вопросов пользователя.',
    )
    add_body(
        document,
        'Задачи проекта: выполнить обзор предметной области и сформулировать '
        'требования; спроектировать структуру интерфейса и поток данных; '
        'реализовать загрузку тем из XML, макет на React, стили SCSS Modules; '
        'обеспечить сохранение переписки в браузере; подготовить раздел '
        'тестирования и руководства пользователя.',
    )
    add_body(
        document,
        'Практическая значимость заключается в демонстрации полного цикла '
        'создания клиентского приложения: от формата данных до интерфейса '
        'и локального хранения состояния.',
    )


def build_main_content(document, args):
    add_section_heading(document, '1 Постановка задачи')

    add_subsection_heading(document, '1.1 Анализ предметной области и аналогов')
    add_body(
        document,
        'Рассмотрены образовательные порталы и справочные разделы '
        'документации (MDN, React). Типичный недостаток статичных страниц — '
        'отсутствие единого сценария «тема → пояснение → чат». '
        'В проекте сочетаются структурированный список тем, раскрывающееся '
        'описание и имитация ответа помощника.',
    )

    add_subsection_heading(document, '1.2 Техническое задание')
    add_body(
        document,
        'Требуется одностраничное приложение с маршрутами главной страницы '
        'и страницы чата. Список тем подгружается из файла data.xml по HTTP. '
        'Пользователь выбирает тему, может раскрыть подробности и перейти '
        'по ссылке. В чате отображаются сообщения пользователя и ответы '
        'помощника (заглушка в коде). Сообщения сохраняются в localStorage.',
    )
    add_body(
        document,
        'Интерфейс должен быть адаптивным: на узком экране список тем '
        'открывается в выезжающей панели. Предусмотреть очистку чата, '
        'экспорт переписки в JSON, копирование последнего ответа.',
    )

    add_subsection_heading(document, '1.3 Выбор средств реализации')
    add_body(
        document,
        'Для вёрстки и логики выбраны HTML5, React 19, React Router, '
        'сборщик Vite, препроцессор SCSS с CSS Modules. Для разбора XML '
        'используются стандартные fetch и DOMParser в браузере. '
        'Хранение сообщений реализовано через Web Storage API.',
    )

    add_subsection_heading(document, '1.4 Выводы по разделу')
    add_body(
        document,
        'Сформулированы требования к функциям приложения и обоснован '
        'стек технологий, позволяющий выполнить задание в среде современного '
        'фронтенд-разработки.',
    )

    document.add_page_break()
    add_section_heading(document, '2 Проектирование веб-приложения')

    add_subsection_heading(document, '2.1 Структура страниц и навигация')
    add_body(
        document,
        'Выделены маршруты «Главная» и «Чат». На главной размещено описание '
        'назначения и инструкция. В чате область разделена на боковую панель '
        'тем и основную колонку с окном сообщений и полем ввода.',
    )

    add_subsection_heading(document, '2.2 Макет интерфейса и компоненты')
    add_body(
        document,
        'Компоненты: Header, Sidebar, ChatWindow, MessageInput. '
        'Семантические элементы header, main, nav, aside облегчают доступность. '
        'Состояние сообщений и выбранной темы сосредоточено в родительском '
        'компоненте страницы чата.',
    )

    add_subsection_heading(document, '2.3 Выводы по разделу')
    add_body(
        document,
        'Спроектирована модульная структура интерфейса и поток данных '
        'от XML к состоянию React.',
    )

    document.add_page_break()
    add_section_heading(document, '3 Реализация')

    add_subsection_heading(document, '3.1 Средства разработки и структура проекта')
    add_body(
        document,
        'Исходный код размещается в каталоге src: страницы в pages, '
        'переиспользуемые части в components, вспомогательные функции в lib, '
        'общие переменные стилей в styles. Статический файл тем лежит в public.',
    )

    add_subsection_heading(document, '3.2 Загрузка и разбор данных в формате XML')
    add_body(
        document,
        'Функция fetchHelpTopics выполняет запрос к data.xml, создаёт '
        'DOMParser и извлекает элементы topic: идентификатор, заголовок, '
        'краткое описание, развёрнутый текст, URL материала. Результат '
        'передаётся в состояние React.',
    )

    add_subsection_heading(document, '3.3 Интерфейс чата и сохранение состояния')
    add_body(
        document,
        'Сообщения добавляются в массив состояния; при ответе помощника '
        'используется асинхронная заглушка с задержкой. Список в области '
        'переписки прокручивается вниз внутри контейнера без смещения всей '
        'страницы браузера. Для экспорта формируется JSON-файл.',
    )

    add_subsection_heading(document, '3.4 Выводы по разделу')
    add_body(
        document,
        'Реализованы ключевые сценарии согласно техническому заданию.',
    )

    document.add_page_break()
    add_section_heading(document, '4 Тестирование и руководство пользователя')

    add_subsection_heading(document, '4.1 Проверка сценариев')
    add_body(
        document,
        'Проверены загрузка XML при доступном сервере разработки, выбор темы, '
        'раскрытие подробностей, отправка сообщений, сохранение после '
        'перезагрузки страницы, очистка чата, экспорт и копирование ответа.',
    )

    add_subsection_heading(document, '4.2 Руководство пользователя')
    add_body(
        document,
        'Пользователь открывает главную страницу, переходит в чат, выбирает '
        'тему слева, при необходимости нажимает «Подробнее», затем вводит '
        'вопрос и отправляет его кнопкой или клавишей Enter.',
    )

    add_subsection_heading(document, '4.3 Выводы по разделу')
    add_body(
        document,
        'Подтверждена работоспособность основных функций приложения.',
    )


def build_conclusion(document):
    document.add_page_break()
    add_center_major(document, 'ЗАКЛЮЧЕНИЕ')
    add_body(
        document,
        'Изучены подходы к организации учебных справочных материалов в вебе.',
    )
    add_body(
        document,
        'Сформулированы требования и выполнено проектирование клиентского '
        'приложения с разделением данных и представления.',
    )
    add_body(
        document,
        'Разработана и испытана программная реализация на React с загрузкой '
        'тем из XML, адаптивным интерфейсом и локальным сохранением переписки.',
    )
    add_body(
        document,
        'Подготовлены материалы для защиты: пояснительная записка и '
        'репозиторий с исходным кодом.',
    )


def build_references(document):
    document.add_page_break()
    add_center_major(document, 'СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ')
    refs = [
        'Документация React [Электронный ресурс]. – Режим доступа: '
        'https://react.dev/ (дата обращения: __________).',
        'Документация Vite [Электронный ресурс]. – Режим доступа: '
        'https://vitejs.dev/ (дата обращения: __________).',
        'MDN Web Docs: DOMParser [Электронный ресурс]. – Режим доступа: '
        'https://developer.mozilla.org/ru/docs/Web/API/DOMParser '
        '(дата обращения: __________).',
        'MDN Web Docs: Web Storage API [Электронный ресурс]. – Режим доступа: '
        'https://developer.mozilla.org/ru/docs/Web/API/Web_Storage_API '
        '(дата обращения: __________).',
    ]
    for i, line in enumerate(refs, start=1):
        p = document.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Mm(12.5)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.left_indent = Mm(0)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(str(i) + '. ' + line)
        set_run_tnr(run, bold=False)


def build_appendix_stub(document):
    document.add_page_break()
    add_center_major(document, 'ПРИЛОЖЕНИЕ А')
    add_body(
        document,
        'Скриншоты главной страницы, страницы чата, раскрытой темы с ссылкой.',
    )
    document.add_page_break()
    add_center_major(document, 'ПРИЛОЖЕНИЕ Б')
    add_body(
        document,
        'Фрагмент файла public/data.xml и пояснение структуры тегов.',
    )
    document.add_page_break()
    add_center_major(document, 'ПРИЛОЖЕНИЕ В')
    add_body(
        document,
        'Листинг ключевого модуля (например, fetchHelpTopics.js или фрагмент '
        'компонента чата) — вставить из IDE с соблюдением полей страницы.',
    )


def parse_args():
    p = argparse.ArgumentParser(
        description='Сгенерировать пояснительную записку (.docx) к курсовому проекту.',
    )
    p.add_argument(
        '-o',
        '--output',
        type=Path,
        default=Path('Пояснительная_записка_курсовой.docx'),
        help='Путь к выходному файлу',
    )
    p.add_argument('--ministry', default='МИНИСТЕРСТВО ОБРАЗОВАНИЯ [страна]')
    p.add_argument('--university', default='Учреждение образования «…»')
    p.add_argument('--faculty', default='_______________')
    p.add_argument('--department', default='_______________')
    p.add_argument('--specialty', default='_______________')
    p.add_argument('--discipline', default='Системный анализ и проектирование')
    p.add_argument(
        '--topic',
        default='Веб-приложение «Интеллектуальный помощник студента»',
    )
    p.add_argument('--student', default='И. О. Фамилия')
    p.add_argument('--group-line', default='студент группы ___')
    p.add_argument('--supervisor', default='И. О. Фамилия')
    p.add_argument(
        '--supervisor-title',
        default='должность, учёная степень',
        dest='supervisor_title',
    )
    return p.parse_args()


def main():
    args = parse_args()
    document = Document()
    configure_section(document)
    configure_normal_style(document)

    build_title_page(document, args)
    document.add_page_break()
    build_assignment_stub(document)
    document.add_page_break()
    build_toc_stub(document)
    document.add_page_break()
    build_intro(document, args)
    document.add_page_break()
    build_main_content(document, args)
    build_conclusion(document)
    build_references(document)
    build_appendix_stub(document)

    out = args.output.resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out))
    print('Сохранено:', out)


if __name__ == '__main__':
    main()
